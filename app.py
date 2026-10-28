import streamlit as st
import pdfplumber
import pandas as pd
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# مكتبات معالجة وإصلاح اللغة العربية القياسية
import arabic_reshaper
from bidi.algorithm import get_display
import re

# PAGE CONFIG
st.set_page_config(page_title="PDF to Excel", page_icon="📂", layout="wide")

st.title("📂 PDF Emails to Excel Converter (Arabic Supported)")
st.write("Upload your Outlook emails PDF file to extract texts and tables correctly.")

# FILE UPLOADER
uploaded_file = st.file_uploader("Choose PDF File", type=["pdf"])

def clean_cell_value(val):
    if val is None:
        return ""
    # إزالة الأسطر والرموز الخاصة والأسهم الحماية لحماية ورقة العمل
    cleaned = re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f\n\r\t\^\*"\'\\\/\[\]\(\)]', ' ', str(val))
    cleaned = cleaned.strip()
    
    # إذا كان النص أطول من 60 حرفاً (جمل التحذير الطويلة للإيميل) يتم تفريغه فوراً حماية للإكسيل
    if len(cleaned) > 60:
        return ""
    return cleaned

def process_pdf(pdf_file):
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    text_font = Font(name="Calibri", size=11, color="333333")
    
    current_row = 1
    
    ws.cell(row=current_row, column=1, value="--- نصوص ومواضيع الإيميلات والجداول المستخرجة ---").font = Font(bold=True, color="FF0000")
    current_row += 2

    # فتح الملف واستخراج المحتوى بالكامل عبر محرك موحد وآمن
    with pdfplumber.open(pdf_file) as pdf:
        for page_num, page in enumerate(pdf.pages, start=1):
            
            # 1. قراءة النصوص العادية بأمان تام داخل try لمنع أي انهيار
            try:
                text = page.extract_text()
                if text:
                    for line in text.split('\n'):
                        if "(cid:" in line:
                            line = re.sub(r'\(cid:\d+\)', '', line)
                        
                        line_cleaned = line.strip()
                        # إذا كان السطر هو النص التحذيري الطويل، نتخطاه ولا نكتبه في الشيت
                        if len(line_cleaned) > 100 or "intended recipient" in line_cleaned.lower():
                            continue
                            
                        if line_cleaned:
                            try:
                                reshaped = arabic_reshaper.reshape(line_cleaned)
                                bidi_line = get_display(reshaped)
                            except:
                                bidi_line = line_cleaned
                                
                            cell = ws.cell(row=current_row, column=1, value=bidi_line)
                            cell.font = text_font
                            cell.alignment = Alignment(horizontal="right")
                            current_row += 1
                    current_row += 1
            except:
                pass # تجاوز أي خطأ تشفير في نصوص الصفحات والانتقال للجداول

            # 2. استخراج الجداول يدوياً بفلتر الحماية الصارم لحجم الخلايا
            try:
                tables = page.extract_tables()
                for table_idx, table in enumerate(tables, start=1):
                    if not table or len(table) == 0:
                        continue
                    
                    ws.cell(row=current_row, column=1, value=f"[Table {table_idx} - Page {page_num}]").font = Font(bold=True)
                    current_row += 1
                    
                    for row_idx, row_data in enumerate(table):
                        has_data = False
                        cleaned_row = []
                        
                        for cell_val in row_data:
                            val = clean_cell_value(cell_val)
                            if val:
                                has_data = True
                                try:
                                    val = get_display(arabic_reshaper.reshape(val))
                                except:
                                    pass
                            cleaned_row.append(val)
                        
                        if not has_data:
                            continue 
                        
                        for col_idx, final_val in enumerate(cleaned_row, start=1):
                            cell = ws.cell(row=current_row, column=col_idx, value=final_val)
                            if row_idx == 0:
                                cell.font = header_font
                                cell.fill = header_fill
                                cell.alignment = Alignment(horizontal="center", vertical="center")
                            else:
                                cell.font = text_font
                                cell.alignment = Alignment(horizontal="right", vertical="center")
                        current_row += 1
                    
                    current_row += 2
            except:
                pass

    # ضبط تلقائي وآمن لعرض الأعمدة في الإكسيل
    for col_idx in range(1, ws.max_column + 1):
        max_len = 0
        col_letter = get_column_letter(col_idx)
        for row_idx in range(1, ws.max_row + 1):
            cell_value = ws.cell(row=row_idx, column=col_idx).value
            if cell_value:
                max_len = max(max_len, len(str(cell_value)))
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)
        
    excel_buffer = BytesIO()
    wb.save(excel_buffer)
    excel_buffer.seek(0)
    return excel_buffer

if uploaded_file is not None:
    with st.spinner("Processing Arabic PDF Safely... Please wait..."):
        try:
            excel_data = process_pdf(uploaded_file)
            st.success("Success!")
            st.download_button(
                label="📥 Download Excel Report",
                data=excel_data,
                file_name="Converted_Arabic_Report.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            )
        except Exception as e:
            st.error(f"Error during processing: {e}")
